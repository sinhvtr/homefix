"""UC12 — đối soát tiền thợ đã thu: tổng hợp, lập/bỏ chọn/chốt/xác nhận/mở lại phiếu, không trùng ThanhToan
giữa hai phiếu, phiếu đã chốt không sửa dòng, số liệu khớp tổng; xuất MISA; màn hình và phân quyền."""

import csv
import io
from datetime import timedelta

import pytest
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.he_thong.models import NhatKyThaoTac
from apps.tai_chinh.dao import DoiSoatDAO
from apps.tai_chinh.models import PhieuDoiSoat, ThanhToan
from apps.tai_chinh.xuat_misa import COT_MISA
from apps.tai_khoan.models import VaiTro
from apps.thong_bao.models import ThongBao
from tests.conftest import NguoiDungFactory, ThoFactory, don_hoan_thanh

pytestmark = pytest.mark.django_db
VS = "VE_SINH_BAO_DUONG"


def _lui_ngay(don, so_ngay):
    ThanhToan.objects.filter(don=don).update(thoi_diem_thu=timezone.now() - timedelta(days=so_ngay))


@pytest.fixture
def ke_toan():
    return NguoiDungFactory(vai_tro=VaiTro.KE_TOAN)


@pytest.fixture
def ba_khoan():
    """Một thợ, ba đơn: tiền mặt 300k; CK cho thợ 200k + vật tư thợ ứng 2 × giá mua 60k; CK công ty + vật tư
    thợ ứng không ghi giá mua (giá bán 50k) + vật tư công ty ứng (không tính)."""
    d1, tho, _ = don_hoan_thanh(VS, so_tien=300_000, hinh_thuc="TIEN_MAT", don_gia=300_000)
    d2, _, _ = don_hoan_thanh(
        VS,
        so_tien=200_000,
        hinh_thuc="CK_TAI_KHOAN_THO",
        tho=tho,
        don_gia=80_000,
        vat_tu=[("Gas R32", 2, 60_000, 60_000, "THO")],
    )
    d3, _, _ = don_hoan_thanh(
        VS,
        hinh_thuc="CK_TAI_KHOAN_CONG_TY",
        tho=tho,
        vat_tu=[("Ống đồng", 1, 50_000, None, "THO"), ("Remote", 1, 90_000, 70_000, "CONG_TY")],
    )
    for don, so_ngay in ((d1, 10), (d2, 6), (d3, 2)):
        _lui_ngay(don, so_ngay)
    return tho, [ThanhToan.objects.get(don=d) for d in (d1, d2, d3)]


# ============================================================== tổng hợp theo thợ


def test_uc12_tong_hop_theo_tho(ke_toan, ba_khoan):
    tho, _ = ba_khoan
    tho_khac = ThoFactory()
    don_hoan_thanh(VS, hinh_thuc="CK_TAI_KHOAN_CONG_TY", tho=tho_khac)  # không còn gì cần đối soát
    nhieu, _, _ = don_hoan_thanh(VS, so_tien=900_000, don_gia=900_000)
    bang = DoiSoatDAO.tong_hop_theo_tho(ke_toan)
    assert [th.tho for th in bang] == [nhieu.tho, tho]  # số phải nộp giảm dần
    th = bang[1]
    assert (th.tien_chua_nop, th.vat_tu_chua_hoan, th.so_phai_nop) == (500_000, 170_000, 330_000)
    assert (th.so_ngay_cu_nhat, th.so_dong_cho_doi_soat, th.ma_tho) == (10, 3, tho.ma_tho)


# ============================================================== lập phiếu — không trùng ThanhToan giữa phiếu


def test_uc12_lap_phieu_gom_moi_khoan_va_so_lieu_khop_tong(ke_toan, ba_khoan):
    tho, ds_tt = ba_khoan
    p = DoiSoatDAO.lap_phieu(ke_toan, tho)
    hom_nay = timezone.localdate()
    assert p.ma == f"DS-{hom_nay:%y%m}-{tho.ma_tho}-1" and p.trang_thai == "NHAP"
    assert (p.tu_ngay, p.den_ngay) == (hom_nay - timedelta(days=10), hom_nay)
    assert set(p.ds_thanh_toan.all()) == set(ds_tt)
    assert (p.tong_tien_thu, p.tong_vat_tu_ung, p.so_phai_nop) == (500_000, 170_000, 330_000)
    dong = DoiSoatDAO.dong_cua_phieu(p)
    assert sum(d.tien_thu for d in dong) == p.tong_tien_thu
    assert sum(d.tien_vat_tu for d in dong) == p.tong_vat_tu_ung
    nk = NhatKyThaoTac.objects.get(hanh_dong="LAP_PHIEU_DOI_SOAT", doi_tuong=f"PhieuDoiSoat:{p.ma}")
    assert nk.nguoi_dung == ke_toan and sorted(nk.du_lieu_sau["thanh_toan"]) == sorted(t.pk for t in ds_tt)


def test_uc12_mot_thanh_toan_khong_the_nam_o_hai_phieu(ke_toan, ba_khoan):
    tho, ds_tt = ba_khoan
    p1 = DoiSoatDAO.lap_phieu(ke_toan, tho)
    with pytest.raises(ValidationError, match="không có khoản nào"):
        DoiSoatDAO.lap_phieu(ke_toan, tho)
    # bỏ chọn một khoản → phiếu thứ hai chỉ lấy đúng khoản đó
    DoiSoatDAO.cap_nhat_dong(ke_toan, p1, [ds_tt[0].pk, ds_tt[2].pk])
    p2 = DoiSoatDAO.lap_phieu(ke_toan, tho)
    assert p2.ma.endswith("-2") and list(p2.ds_thanh_toan.all()) == [ds_tt[1]]
    # cố chọn lại vào p1 một khoản đã ở p2 → từ chối, không đổi gì
    with pytest.raises(ValidationError, match="phiếu khác"):
        DoiSoatDAO.cap_nhat_dong(ke_toan, p1, [t.pk for t in ds_tt])
    assert ThanhToan.objects.get(pk=ds_tt[1].pk).phieu_doi_soat == p2
    p1.refresh_from_db()
    assert (p1.so_phai_nop, p2.so_phai_nop) == (250_000, 80_000)  # tổng hai phiếu = tổng phải nộp của thợ
    # cấu trúc CSDL: ThanhToan có đúng một cột phieu_doi_soat → không thể thuộc hai phiếu
    assert ThanhToan._meta.get_field("phieu_doi_soat").many_to_one


def test_uc12_bo_chon_cap_nhat_tong_va_ghi_nhat_ky(ke_toan, ba_khoan):
    tho, ds_tt = ba_khoan
    p = DoiSoatDAO.lap_phieu(ke_toan, tho)
    p = DoiSoatDAO.cap_nhat_dong(ke_toan, p, [ds_tt[0].pk])
    assert (p.tong_tien_thu, p.tong_vat_tu_ung, p.so_phai_nop) == (300_000, 0, 300_000)
    assert ThanhToan.objects.filter(pk__in=[ds_tt[1].pk, ds_tt[2].pk], phieu_doi_soat=None).count() == 2
    assert [d.thanh_toan for d in DoiSoatDAO.dong_co_the_them(p)] == [ds_tt[1], ds_tt[2]]
    p = DoiSoatDAO.cap_nhat_dong(ke_toan, p, [t.pk for t in ds_tt])  # chọn lại
    assert p.so_phai_nop == 330_000
    assert NhatKyThaoTac.objects.filter(hanh_dong="SUA_DONG_DOI_SOAT").count() == 2


def test_uc12_khong_chon_khoan_cua_tho_khac(ke_toan, ba_khoan):
    tho, _ = ba_khoan
    p = DoiSoatDAO.lap_phieu(ke_toan, tho)
    khac, _, _ = don_hoan_thanh(VS, so_tien=100_000)
    with pytest.raises(ValidationError):
        DoiSoatDAO.cap_nhat_dong(
            ke_toan, p, [*p.ds_thanh_toan.values_list("pk", flat=True), ThanhToan.objects.get(don=khac).pk]
        )


# ============================================================== chốt, xác nhận đã nộp, mở lại


def test_uc12_phieu_da_chot_khong_sua_dong_roi_xac_nhan_nop(ke_toan, ba_khoan):
    tho, ds_tt = ba_khoan
    p = DoiSoatDAO.chot(ke_toan, DoiSoatDAO.lap_phieu(ke_toan, tho))
    assert p.trang_thai == "DA_CHOT"
    assert ThongBao.objects.filter(nguoi_nhan=tho.nguoi_dung, tieu_de__contains=p.ma).exists()
    with pytest.raises(ValidationError, match="Đã chốt"):
        DoiSoatDAO.cap_nhat_dong(ke_toan, p, [ds_tt[0].pk])
    with pytest.raises(ValidationError):
        DoiSoatDAO.huy_nhap(ke_toan, p)
    assert p.ds_thanh_toan.count() == 3
    p = DoiSoatDAO.xac_nhan_da_nop(ke_toan, p)
    assert p.trang_thai == "DA_NOP" and p.xac_nhan_luc is not None
    assert [ThanhToan.objects.get(pk=t.pk).trang_thai for t in ds_tt] == ["DA_NOP", "DA_NOP", "KHONG_CAN_NOP"]
    assert DoiSoatDAO.tong_hop_theo_tho(ke_toan) == []  # đã nộp, vật tư đã hoàn
    assert NhatKyThaoTac.objects.filter(
        hanh_dong="XAC_NHAN_NOP_TIEN", doi_tuong=f"PhieuDoiSoat:{p.ma}"
    ).exists()
    with pytest.raises(ValidationError):
        DoiSoatDAO.xac_nhan_da_nop(ke_toan, p)  # không xác nhận hai lần


def test_uc12_mo_lai_phieu_can_ly_do_va_tra_ve_chua_nop(ke_toan, ba_khoan):
    tho, ds_tt = ba_khoan
    p = DoiSoatDAO.xac_nhan_da_nop(ke_toan, DoiSoatDAO.chot(ke_toan, DoiSoatDAO.lap_phieu(ke_toan, tho)))
    with pytest.raises(ValidationError):
        DoiSoatDAO.mo_lai(ke_toan, p, " ")
    p = DoiSoatDAO.mo_lai(ke_toan, p, "Nhầm số tiền đơn 2")
    assert (p.trang_thai, p.xac_nhan_luc) == ("NHAP", None)
    assert ThanhToan.objects.get(pk=ds_tt[0].pk).trang_thai == "CHUA_NOP"
    assert ThanhToan.objects.get(pk=ds_tt[2].pk).trang_thai == "KHONG_CAN_NOP"
    nk = NhatKyThaoTac.objects.get(hanh_dong="MO_LAI_PHIEU_DOI_SOAT")
    assert nk.du_lieu_truoc["trang_thai"] == "DA_NOP" and nk.du_lieu_sau["ly_do"] == "Nhầm số tiền đơn 2"
    DoiSoatDAO.cap_nhat_dong(ke_toan, p, [ds_tt[0].pk])  # nháp lại → sửa được


def test_uc12_so_phai_nop_am_khi_vat_tu_ung_lon_hon_tien_thu(ke_toan):
    don, tho, _ = don_hoan_thanh(
        VS, hinh_thuc="CK_TAI_KHOAN_CONG_TY", vat_tu=[("Block máy nén", 1, 2_500_000, 1_800_000, "THO")]
    )
    p = DoiSoatDAO.lap_phieu(ke_toan, tho)
    assert (p.tong_tien_thu, p.tong_vat_tu_ung, p.so_phai_nop) == (0, 1_800_000, -1_800_000)


def test_uc12_rang_buoc_csdl_so_phai_nop_khop(ke_toan, ba_khoan):
    tho, _ = ba_khoan
    p = DoiSoatDAO.lap_phieu(ke_toan, tho)
    with pytest.raises(IntegrityError), transaction.atomic():
        PhieuDoiSoat.objects.filter(pk=p.pk).update(so_phai_nop=1)


def test_uc12_ky_mac_dinh_tu_sau_phieu_gan_nhat(ke_toan, ba_khoan):
    tho, ds_tt = ba_khoan
    hom_nay = timezone.localdate()
    p1 = DoiSoatDAO.lap_phieu(ke_toan, tho, den_ngay=hom_nay - timedelta(days=5))
    assert set(p1.ds_thanh_toan.all()) == {ds_tt[0], ds_tt[1]}  # đơn 2 ngày trước nằm ngoài kỳ
    p2 = DoiSoatDAO.lap_phieu(ke_toan, tho)
    assert (p2.tu_ngay, p2.den_ngay) == (hom_nay - timedelta(days=4), hom_nay)
    assert list(p2.ds_thanh_toan.all()) == [ds_tt[2]]


def test_uc12_chi_ke_toan_ghi(ba_khoan):
    tho, _ = ba_khoan
    for vai_tro in (VaiTro.QUAN_TRI, VaiTro.TONG_DAI, VaiTro.GIAM_DOC):
        with pytest.raises(PermissionDenied):
            DoiSoatDAO.lap_phieu(NguoiDungFactory(vai_tro=vai_tro), tho)
    with pytest.raises(PermissionDenied):
        DoiSoatDAO.tong_hop_theo_tho(NguoiDungFactory(vai_tro=VaiTro.TONG_DAI))


# ============================================================== xuất MISA


def test_uc12_xuat_misa_csv_xlsx(client, ke_toan, ba_khoan):
    tho, _ = ba_khoan
    p = DoiSoatDAO.xac_nhan_da_nop(ke_toan, DoiSoatDAO.chot(ke_toan, DoiSoatDAO.lap_phieu(ke_toan, tho)))
    DoiSoatDAO.chot(
        ke_toan,
        DoiSoatDAO.lap_phieu(ke_toan, ThanhToan.objects.get(don=don_hoan_thanh(VS, so_tien=1)[0]).tho),
    )  # chưa nộp → không xuất
    client.force_login(ke_toan)
    hom_nay = timezone.localdate()
    phan_hoi = client.get(f"/ke-toan/xuat/?dinh_dang=csv&tu={hom_nay}&den={hom_nay}")
    assert phan_hoi["Content-Type"].startswith("text/csv") and "attachment" in phan_hoi["Content-Disposition"]
    noi_dung = phan_hoi.content.decode("utf-8")
    assert noi_dung.startswith("﻿")
    dong = list(csv.reader(io.StringIO(noi_dung.lstrip("﻿"))))
    assert dong[0] == [ten for ten, _ in COT_MISA]
    assert dong[1:] == [[f"{hom_nay:%d/%m/%Y}", p.ma, dong[1][2], tho.ma_tho, tho.ho_ten, "330000", "", ""]]
    assert tho.ho_ten in dong[1][2]
    wb = load_workbook(
        io.BytesIO(client.get(f"/ke-toan/xuat/?dinh_dang=xlsx&tu={hom_nay}&den={hom_nay}").content)
    )
    ws = wb.active
    assert [o.value for o in ws[1]][:2] == ["Ngày chứng từ", "Số chứng từ"] and ws["F2"].value == 330_000
    assert ws.max_row == 2
    # ngoài kỳ → chỉ có dòng tiêu đề
    hom_qua = hom_nay - timedelta(days=1)
    assert (
        len(
            client.get(f"/ke-toan/xuat/?dinh_dang=csv&tu={hom_qua}&den={hom_qua}")
            .content.decode()
            .splitlines()
        )
        == 1
    )


# ============================================================== màn hình


def test_man_hinh_ke_toan_lap_chot_xac_nhan(client, ke_toan, ba_khoan):
    tho, ds_tt = ba_khoan
    client.force_login(ke_toan)
    trang = client.get("/ke-toan/doi-soat/").content.decode()
    assert tho.ho_ten in trang and "330.000 đ" in trang and "Lập phiếu (3 khoản)" in trang
    phan_hoi = client.post("/ke-toan/doi-soat/lap/", {"tho": tho.pk})
    p = PhieuDoiSoat.objects.get()
    assert phan_hoi.url == f"/ke-toan/doi-soat/{p.ma}/"
    chi_tiet = client.get(phan_hoi.url).content.decode()
    assert "Chốt phiếu" in chi_tiet and ds_tt[0].don.ma_don in chi_tiet
    client.post(phan_hoi.url, {"hanh_dong": "cap_nhat", "thanh_toan": [t.pk for t in ds_tt]})
    # bỏ tích khoản thứ ba rồi bấm thẳng “Chốt” → chốt đúng các khoản đang tích
    client.post(phan_hoi.url, {"hanh_dong": "chot", "thanh_toan": [ds_tt[0].pk, ds_tt[1].pk]})
    p.refresh_from_db()
    assert (p.trang_thai, p.so_phai_nop) == ("DA_CHOT", 380_000)
    assert "Chốt phiếu" not in client.get(phan_hoi.url).content.decode()
    client.post(phan_hoi.url, {"hanh_dong": "da_nop"})
    p.refresh_from_db()
    assert p.trang_thai == "DA_NOP"


@pytest.mark.parametrize("vai_tro", [VaiTro.QUAN_TRI, VaiTro.GIAM_DOC])
def test_man_hinh_quan_tri_giam_doc_chi_xem(client, ke_toan, ba_khoan, vai_tro):
    tho, _ = ba_khoan
    p = DoiSoatDAO.lap_phieu(ke_toan, tho)
    client.force_login(NguoiDungFactory(vai_tro=vai_tro))
    trang = client.get("/ke-toan/doi-soat/").content.decode()
    assert p.ma in trang and "Lập phiếu" not in trang
    chi_tiet = client.get(f"/ke-toan/doi-soat/{p.ma}/")
    assert chi_tiet.status_code == 200 and "Chốt phiếu" not in chi_tiet.content.decode()
    assert client.post(f"/ke-toan/doi-soat/{p.ma}/", {"hanh_dong": "chot"}).status_code == 403
    assert client.post("/ke-toan/doi-soat/lap/", {"tho": tho.pk}).status_code == 403


def test_man_hinh_tho_chi_xem_phieu_cua_minh(client, ke_toan, ba_khoan):
    tho, _ = ba_khoan
    p = DoiSoatDAO.chot(ke_toan, DoiSoatDAO.lap_phieu(ke_toan, tho))
    client.force_login(tho.nguoi_dung)
    trang = client.get("/tho/doi-soat/").content.decode()
    assert p.ma in trang and "Phải nộp 330.000 đ" in trang
    assert client.get(f"/ke-toan/doi-soat/{p.ma}/").status_code == 403
    tho_khac = ThoFactory()
    client.force_login(tho_khac.nguoi_dung)
    assert p.ma not in client.get("/tho/doi-soat/").content.decode()
    with pytest.raises(PhieuDoiSoat.DoesNotExist):
        DoiSoatDAO.lay_phieu(tho_khac.nguoi_dung, p.ma)


def test_ke_toan_dang_nhap_vao_trang_doi_soat(client, ke_toan):
    from tests.conftest import MAT_KHAU_THU

    phan_hoi = client.post(
        "/dang-nhap/", {"username": ke_toan.username, "password": MAT_KHAU_THU}, follow=True
    )
    assert phan_hoi.redirect_chain[-1][0] == "/ke-toan/doi-soat/" and phan_hoi.status_code == 200
