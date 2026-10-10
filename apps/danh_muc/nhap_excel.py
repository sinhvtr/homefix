"""Nhập thợ và bảng giá hàng loạt từ Excel (UC13, SHOULD).

Quy tắc (A-32): kiểm tra toàn bộ tệp trước; có lỗi ở bất kỳ dòng nào thì KHÔNG ghi dòng nào và báo lỗi theo số dòng.
"""

import io
import unicodedata
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from apps.he_thong.so_dien_thoai import chuan_hoa_sdt
from apps.tai_khoan.models import NguoiDung
from apps.tai_khoan.phan_quyen import CHI_QUAN_TRI, kiem_tra_vai_tro

from .dao import DichVuDAO, ThoDAO
from .models import DichVu, KhuVuc, LoaiThietBi, Tho

COT_THO = [
    "Họ tên*",
    "Số điện thoại*",
    "Hình thức*",
    "Kỹ năng (mã dịch vụ, cách nhau dấu phẩy)",
    "Khu vực phụ trách (mã khu vực, cách nhau dấu phẩy)",
    "Ghi chú",
]
COT_DICH_VU = [
    "Mã*",
    "Tên dịch vụ*",
    "Nhóm*",
    "Loại thiết bị*",
    "Đơn giá tham khảo (đ)*",
    "Đơn vị tính*",
    "Bảo hành (ngày)",
    "Đang áp dụng (Có/Không)",
]
DONG_TOI_DA = 2000


@dataclass
class KetQuaNhap:
    so_tao: int = 0
    so_cap_nhat: int = 0
    loi: list[tuple[int, str]] = field(default_factory=list)
    ds_pin: list[tuple[str, str, str]] = field(default_factory=list)  # (họ tên, SĐT, PIN) — hiện một lần

    @property
    def thanh_cong(self) -> bool:
        return not self.loi


def _bo_dau(chuoi: str) -> str:
    chuoi = unicodedata.normalize("NFD", chuoi.replace("đ", "d").replace("Đ", "D"))
    return "".join(c for c in chuoi if unicodedata.category(c) != "Mn")


def _ma_lua_chon(gia_tri, choices, ten_cot: str) -> str:
    """Nhận mã (VE_SINH_BAO_DUONG) hoặc nhãn (Vệ sinh, bảo dưỡng), không phân biệt hoa thường/dấu."""
    chuoi = _bo_dau(str(gia_tri or "").strip()).upper().replace(" ", "_").replace(",", "")
    for ma, nhan in choices:
        if chuoi in (ma, _bo_dau(nhan).upper().replace(" ", "_").replace(",", "")):
            return ma
    hop_le = ", ".join(ma for ma, _ in choices)
    raise ValidationError(f"{ten_cot} “{gia_tri}” không hợp lệ (dùng một trong: {hop_le})")


def _so_nguyen(gia_tri, ten_cot: str, bat_buoc=True) -> int | None:
    if gia_tri in (None, ""):
        if bat_buoc:
            raise ValidationError(f"Thiếu {ten_cot}")
        return None
    if isinstance(gia_tri, bool):
        raise ValidationError(f"{ten_cot} phải là số")
    if isinstance(gia_tri, int):
        so = gia_tri
    elif isinstance(gia_tri, float):
        if not gia_tri.is_integer():
            raise ValidationError(f"{ten_cot} phải là số nguyên")
        so = int(gia_tri)
    else:
        chuoi = (
            str(gia_tri).strip().lower().replace("đ", "").replace(".", "").replace(" ", "").replace(",", "")
        )
        if not chuoi.isdigit():
            raise ValidationError(f"{ten_cot} “{gia_tri}” không phải số nguyên")
        so = int(chuoi)
    if so < 0:
        raise ValidationError(f"{ten_cot} không được âm")
    return so


def _chuoi(gia_tri) -> str:
    if gia_tri is None:
        return ""
    if isinstance(gia_tri, float) and gia_tri.is_integer():
        gia_tri = int(gia_tri)
    return str(gia_tri).strip()


def _sdt_tu_o(gia_tri) -> str:
    """Ô kiểu số trong Excel làm mất số 0 đầu (0977… → 977…): thêm lại khi còn đúng 9 chữ số."""
    chuoi = _chuoi(gia_tri)
    if isinstance(gia_tri, int | float) and len(chuoi) == 9 and chuoi.isdigit():
        return "0" + chuoi
    return chuoi


def _danh_sach_ma(gia_tri) -> list[str]:
    return [m.strip() for m in _chuoi(gia_tri).replace(";", ",").split(",") if m.strip()]


def _doc_dong(tep, so_cot: int):
    """Sinh (số dòng Excel, danh sách ô) cho các dòng có dữ liệu, bỏ dòng tiêu đề."""
    try:
        wb = load_workbook(tep, read_only=True, data_only=True)
    except Exception as loi:
        raise ValidationError("Không đọc được tệp Excel. Hãy dùng tệp mẫu tải về từ hệ thống.") from loi
    ws = wb.worksheets[0]
    for so_dong, dong in enumerate(ws.iter_rows(min_row=2, max_col=so_cot, values_only=True), start=2):
        if so_dong > DONG_TOI_DA + 1:
            raise ValidationError(f"Tệp quá {DONG_TOI_DA} dòng.")
        o = list(dong) + [None] * (so_cot - len(dong))
        if all(v in (None, "") for v in o):
            continue
        yield so_dong, o


def _tao_mau(cot: list[str], huong_dan: list[tuple[str, str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Dữ liệu"
    ws.append(cot)
    for i, _ in enumerate(cot, start=1):
        ws.cell(row=1, column=i).font = Font(bold=True)
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 28
    hd = wb.create_sheet("Hướng dẫn")
    hd.append(["Mục", "Giá trị hợp lệ"])
    hd["A1"].font = hd["B1"].font = Font(bold=True)
    for muc, gia_tri in huong_dan:
        hd.append([muc, gia_tri])
    hd.column_dimensions["A"].width = 30
    hd.column_dimensions["B"].width = 100
    dem = io.BytesIO()
    wb.save(dem)
    return dem.getvalue()


def _lua_chon(choices) -> str:
    return "; ".join(f"{ma} ({nhan})" for ma, nhan in choices)


# ------------------------------------------------------------------ Thợ


def tao_mau_tho() -> bytes:
    return _tao_mau(
        COT_THO,
        [
            ("Dòng 1", "Tiêu đề — giữ nguyên. Điền từ dòng 2 trên trang “Dữ liệu”. Cột có * là bắt buộc."),
            ("Hình thức", _lua_chon(Tho.HinhThuc.choices)),
            ("Mã dịch vụ", ", ".join(DichVu.objects.order_by("ma").values_list("ma", flat=True))),
            ("Mã khu vực", ", ".join(KhuVuc.objects.order_by("ma").values_list("ma", flat=True))),
            ("PIN", "Hệ thống tự cấp PIN 6 số cho từng thợ, hiện một lần sau khi nhập."),
        ],
    )


def nhap_tho(nguoi_dung, tep) -> KetQuaNhap:
    kiem_tra_vai_tro(nguoi_dung, *CHI_QUAN_TRI)
    kq = KetQuaNhap()
    dich_vu = {dv.ma.upper(): dv for dv in DichVu.objects.all()}
    khu_vuc = {kv.ma.upper(): kv for kv in KhuVuc.objects.all()}
    sdt_da_co = set(NguoiDung.objects.values_list("so_dien_thoai", flat=True))
    sdt_trong_tep: dict[str, int] = {}
    hop_le = []
    try:
        dong_doc = list(_doc_dong(tep, len(COT_THO)))
    except ValidationError as loi:
        kq.loi.append((0, " ".join(loi.messages)))
        return kq
    for so_dong, (ho_ten, sdt, hinh_thuc, ky_nang, kv, ghi_chu) in dong_doc:
        loi_dong = []
        ho_ten = _chuoi(ho_ten)
        if not ho_ten:
            loi_dong.append("Thiếu họ tên")
        try:
            sdt = chuan_hoa_sdt(_sdt_tu_o(sdt))
            if sdt in sdt_da_co:
                loi_dong.append(f"Số điện thoại {sdt} đã có tài khoản")
            elif sdt in sdt_trong_tep:
                loi_dong.append(f"Số điện thoại {sdt} trùng với dòng {sdt_trong_tep[sdt]}")
            sdt_trong_tep.setdefault(sdt, so_dong)
        except ValidationError as loi:
            loi_dong.extend(loi.messages)
        try:
            hinh_thuc = _ma_lua_chon(hinh_thuc, Tho.HinhThuc.choices, "Hình thức")
        except ValidationError as loi:
            loi_dong.extend(loi.messages)
        ma_dv = _danh_sach_ma(ky_nang)
        thieu = [m for m in ma_dv if m.upper() not in dich_vu]
        if thieu:
            loi_dong.append(f"Không có dịch vụ mã: {', '.join(thieu)}")
        ma_kv = _danh_sach_ma(kv)
        thieu = [m for m in ma_kv if m.upper() not in khu_vuc]
        if thieu:
            loi_dong.append(f"Không có khu vực mã: {', '.join(thieu)}")
        if loi_dong:
            kq.loi.append((so_dong, "; ".join(loi_dong)))
            continue
        hop_le.append(
            {
                "ho_ten": ho_ten,
                "so_dien_thoai": sdt,
                "hinh_thuc": hinh_thuc,
                "ky_nang": [dich_vu[m.upper()] for m in ma_dv],
                "khu_vuc_phu_trach": [khu_vuc[m.upper()] for m in ma_kv],
                "ghi_chu": _chuoi(ghi_chu),
            }
        )
    if not dong_doc:
        kq.loi.append((0, "Tệp không có dòng dữ liệu nào."))
    if kq.loi:
        return kq
    with transaction.atomic():
        for du_lieu in hop_le:
            tho, pin = ThoDAO.tao(nguoi_dung, du_lieu)
            kq.so_tao += 1
            kq.ds_pin.append((tho.ho_ten, tho.so_dien_thoai, pin))
    return kq


# ------------------------------------------------------------------ Bảng giá


def tao_mau_bang_gia() -> bytes:
    return _tao_mau(
        COT_DICH_VU,
        [
            ("Dòng 1", "Tiêu đề — giữ nguyên. Điền từ dòng 2 trên trang “Dữ liệu”. Cột có * là bắt buộc."),
            ("Mã", "Mã đã có → cập nhật dịch vụ đó (đổi giá được ghi nhật ký); mã mới → thêm dịch vụ."),
            ("Nhóm", _lua_chon(DichVu.NhomDichVu.choices)),
            ("Loại thiết bị", _lua_chon(LoaiThietBi.choices)),
            ("Đơn giá", "Số nguyên, đơn vị đồng, ví dụ 150000."),
            ("Bảo hành (ngày)", "Để trống = mặc định theo nhóm (7 / 90 ngày)."),
            ("Đang áp dụng", "Có hoặc Không; để trống = Có."),
        ],
    )


def nhap_bang_gia(nguoi_dung, tep) -> KetQuaNhap:
    kiem_tra_vai_tro(nguoi_dung, *CHI_QUAN_TRI)
    kq = KetQuaNhap()
    hien_co = {dv.ma: dv for dv in DichVu.objects.all()}
    ma_trong_tep: dict[str, int] = {}
    hop_le = []
    try:
        dong_doc = list(_doc_dong(tep, len(COT_DICH_VU)))
    except ValidationError as loi:
        kq.loi.append((0, " ".join(loi.messages)))
        return kq
    for so_dong, (ma, ten, nhom, loai, don_gia, dvt, bao_hanh, ap_dung) in dong_doc:
        loi_dong = []
        du_lieu = {}
        ma = _chuoi(ma).upper()
        if not ma:
            loi_dong.append("Thiếu mã")
        elif len(ma) > 20:
            loi_dong.append("Mã dài quá 20 ký tự")
        elif ma in ma_trong_tep:
            loi_dong.append(f"Mã {ma} trùng với dòng {ma_trong_tep[ma]}")
        ma_trong_tep.setdefault(ma, so_dong)
        du_lieu["ma"] = ma
        du_lieu["ten_dich_vu"] = _chuoi(ten)
        if not du_lieu["ten_dich_vu"]:
            loi_dong.append("Thiếu tên dịch vụ")
        du_lieu["don_vi_tinh"] = _chuoi(dvt)
        if not du_lieu["don_vi_tinh"]:
            loi_dong.append("Thiếu đơn vị tính")
        for ten_truong, gia_tri, choices, ten_cot in (
            ("nhom_dich_vu", nhom, DichVu.NhomDichVu.choices, "Nhóm"),
            ("loai_thiet_bi", loai, LoaiThietBi.choices, "Loại thiết bị"),
        ):
            try:
                du_lieu[ten_truong] = _ma_lua_chon(gia_tri, choices, ten_cot)
            except ValidationError as loi:
                loi_dong.extend(loi.messages)
        try:
            du_lieu["don_gia_tham_khao"] = _so_nguyen(don_gia, "đơn giá")
            du_lieu["thoi_han_bao_hanh_ngay"] = _so_nguyen(bao_hanh, "bảo hành", bat_buoc=False)
        except ValidationError as loi:
            loi_dong.extend(loi.messages)
        ap_dung = _bo_dau(_chuoi(ap_dung)).lower()
        if ap_dung in ("", "co", "x", "1", "true"):
            du_lieu["dang_ap_dung"] = True
        elif ap_dung in ("khong", "0", "false"):
            du_lieu["dang_ap_dung"] = False
        else:
            loi_dong.append("Đang áp dụng phải là Có hoặc Không")
        if loi_dong:
            kq.loi.append((so_dong, "; ".join(loi_dong)))
        else:
            hop_le.append(du_lieu)
    if not dong_doc:
        kq.loi.append((0, "Tệp không có dòng dữ liệu nào."))
    if kq.loi:
        return kq
    with transaction.atomic():
        for du_lieu in hop_le:
            if du_lieu["ma"] in hien_co:
                DichVuDAO.sua(nguoi_dung, hien_co[du_lieu["ma"]], du_lieu)
                kq.so_cap_nhat += 1
            else:
                DichVuDAO.tao(nguoi_dung, du_lieu)
                kq.so_tao += 1
    return kq
