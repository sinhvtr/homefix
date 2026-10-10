"""Nhập thợ và bảng giá từ Excel: tải mẫu → điền → tải lên, báo lỗi đúng dòng (UC13)."""

import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.danh_muc.models import DichVu, Tho
from apps.danh_muc.nhap_excel import COT_DICH_VU, COT_THO
from apps.he_thong.models import NhatKyThaoTac
from apps.tai_khoan.models import NguoiDung, VaiTro
from tests.conftest import DichVuFactory, KhuVucFactory

pytestmark = pytest.mark.django_db

LOAI = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _tep(cot, dong) -> SimpleUploadedFile:
    wb = Workbook()
    ws = wb.active
    ws.append(cot)
    for d in dong:
        ws.append(d)
    dem = io.BytesIO()
    wb.save(dem)
    return SimpleUploadedFile("du_lieu.xlsx", dem.getvalue(), content_type=LOAI)


@pytest.mark.parametrize(
    ("url", "cot"), [("/danh-muc/tho/mau.xlsx", COT_THO), ("/danh-muc/dich-vu/mau.xlsx", COT_DICH_VU)]
)
def test_tai_tep_mau(client_quan_tri, url, cot):
    DichVuFactory(ma="VS-DH-TT")
    phan_hoi = client_quan_tri.get(url)
    assert phan_hoi.status_code == 200 and phan_hoi["Content-Type"] == LOAI
    wb = load_workbook(io.BytesIO(phan_hoi.content))
    assert [c.value for c in wb.worksheets[0][1]] == cot
    assert wb.sheetnames == ["Dữ liệu", "Hướng dẫn"]


# ------------------------------------------------------------ Thợ


def test_nhap_tho_thanh_cong_hien_pin_mot_lan(client_quan_tri):
    DichVuFactory(ma="VS-DH-TT")
    DichVuFactory(ma="SC-DH-GAS")
    KhuVucFactory(ma="CAU_GIAY")
    tep = _tep(
        COT_THO,
        [
            ["Nguyễn Văn A", "0977 000 001", "Chính thức", "VS-DH-TT, sc-dh-gas", "CAU_GIAY", ""],
            ["Lê Thị B", 977000002, "CONG_TAC_VIEN", "", "", "Mới vào"],
        ],
    )
    phan_hoi = client_quan_tri.post("/danh-muc/tho/nhap/", {"tep": tep})
    noi_dung = phan_hoi.content.decode()
    assert phan_hoi.status_code == 200
    assert Tho.objects.count() == 2
    a = Tho.objects.get(so_dien_thoai="0977000001")
    assert a.hinh_thuc == "CHINH_THUC" and a.ky_nang.count() == 2 and a.khu_vuc_phu_trach.count() == 1
    assert Tho.objects.get(so_dien_thoai="0977000002").hinh_thuc == "CONG_TAC_VIEN"
    assert "PIN của thợ vừa nhập" in noi_dung and "0977000001" in noi_dung
    assert NguoiDung.objects.filter(vai_tro=VaiTro.THO, phai_doi_mat_khau=True).count() == 2


def test_nhap_tho_bao_loi_dung_dong_va_khong_ghi_gi(client_quan_tri, tao_nguoi_dung):
    tao_nguoi_dung(so_dien_thoai="0977000009")
    tep = _tep(
        COT_THO,
        [
            ["Hợp lệ", "0977000001", "CHINH_THUC", "", "", ""],  # dòng 2
            ["", "12345", "Thợ cả", "KHONG_CO", "KV_LA", ""],  # dòng 3
            [None, None, None, None, None, None],  # dòng 4 trống — bỏ qua
            ["Trùng", "0977000001", "CHINH_THUC", "", "", ""],  # dòng 5
            ["Đã có", "0977000009", "CHINH_THUC", "", "", ""],  # dòng 6
        ],
    )
    noi_dung = client_quan_tri.post("/danh-muc/tho/nhap/", {"tep": tep}).content.decode()
    assert not Tho.objects.exists()
    assert "Chưa nhập — 3 dòng lỗi" in noi_dung
    for mau in ["Thiếu họ tên", "Số điện thoại không hợp lệ", "Hình thức “Thợ cả” không hợp lệ",
                "Không có dịch vụ mã: KHONG_CO", "Không có khu vực mã: KV_LA",
                "trùng với dòng 2", "0977000009 đã có tài khoản"]:  # fmt: skip
        assert mau in noi_dung, mau
    assert '<td class="p-1 font-mono">3</td>' in noi_dung
    assert '<td class="p-1 font-mono">5</td>' in noi_dung
    assert '<td class="p-1 font-mono">6</td>' in noi_dung


def test_nhap_tep_khong_phai_xlsx(client_quan_tri):
    tep = SimpleUploadedFile("tho.csv", b"a,b", content_type="text/csv")
    phan_hoi = client_quan_tri.post("/danh-muc/tho/nhap/", {"tep": tep})
    assert "Chỉ nhận tệp .xlsx" in phan_hoi.content.decode()


def test_nhap_tep_hong(client_quan_tri):
    tep = SimpleUploadedFile("tho.xlsx", b"khong phai excel", content_type=LOAI)
    phan_hoi = client_quan_tri.post("/danh-muc/tho/nhap/", {"tep": tep})
    assert "Không đọc được tệp Excel" in phan_hoi.content.decode()


# ------------------------------------------------------------ Bảng giá


def test_nhap_bang_gia_them_moi_va_cap_nhat_ghi_nhat_ky_gia(client_quan_tri, quan_tri):
    DichVuFactory(ma="VS-DH-TT", don_gia_tham_khao=150_000)
    tep = _tep(
        COT_DICH_VU,
        [
            [
                "VS-DH-TT",
                "Vệ sinh điều hòa treo tường",
                "VE_SINH_BAO_DUONG",
                "Điều hòa",
                160000.0,
                "máy",
                None,
                "Có",
            ],
            ["sc-dh-gas", "Nạp gas điều hòa", "Sửa chữa, thay vật tư", "DIEU_HOA", "350.000", "máy", "", ""],
            ["VS-TL", "Vệ sinh tủ lạnh", "VE_SINH_BAO_DUONG", "TU_LANH", 200000, "máy", 14, "Không"],
        ],
    )
    phan_hoi = client_quan_tri.post("/danh-muc/dich-vu/nhap/", {"tep": tep})
    assert phan_hoi.status_code == 302
    assert DichVu.objects.get(ma="VS-DH-TT").don_gia_tham_khao == 160_000
    gas = DichVu.objects.get(ma="SC-DH-GAS")
    assert (gas.don_gia_tham_khao, gas.thoi_han_bao_hanh_ngay, gas.nhom_dich_vu) == (
        350_000,
        90,
        "SUA_CHUA_THAY_VAT_TU",
    )
    tl = DichVu.objects.get(ma="VS-TL")
    assert (tl.thoi_han_bao_hanh_ngay, tl.dang_ap_dung) == (14, False)
    nk = NhatKyThaoTac.objects.get(hanh_dong="SUA_GIA")
    assert (nk.nguoi_dung, nk.du_lieu_truoc, nk.du_lieu_sau) == (
        quan_tri,
        {"don_gia_tham_khao": 150_000},
        {"don_gia_tham_khao": 160_000},
    )


def test_nhap_bang_gia_loi_dong(client_quan_tri):
    tep = _tep(
        COT_DICH_VU,
        [
            ["OK-1", "Hợp lệ", "VE_SINH_BAO_DUONG", "DIEU_HOA", 100000, "máy", None, None],
            ["", "Thiếu mã", "SAI_NHOM", "DIEU_HOA", 150000.5, "máy", -1, "Có lẽ"],
        ],
    )
    noi_dung = client_quan_tri.post("/danh-muc/dich-vu/nhap/", {"tep": tep}).content.decode()
    assert not DichVu.objects.exists()
    for mau in [
        "Thiếu mã",
        "Nhóm “SAI_NHOM” không hợp lệ",
        "đơn giá phải là số nguyên",
        "Đang áp dụng phải là Có hoặc Không",
    ]:
        assert mau in noi_dung, mau
    assert '<td class="p-1 font-mono">3</td>' in noi_dung
