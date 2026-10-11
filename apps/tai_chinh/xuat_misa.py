"""Xuất dữ liệu đối soát cho MISA (UC12, OI-07 — chưa có mẫu cột MISA thật).

Định nghĩa cột đặt trong COT_MISA: sửa tiêu đề / thứ tự / cách lấy giá trị ở đây khi có mẫu thật, không cần sửa nơi
khác. Mỗi dòng = một phiếu đối soát đã xác nhận nộp (DA_NOP) có ngày xác nhận trong kỳ (A-81).
"""

import csv
import io
from datetime import date

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font

from apps.tai_khoan.phan_quyen import kiem_tra_vai_tro

from .dao import XEM_DOI_SOAT
from .models import PhieuDoiSoat


def _dien_giai(p: PhieuDoiSoat) -> str:
    if p.so_phai_nop >= 0:
        return f"Thu tiền thợ {p.tho.ho_ten} nộp, đối soát {p.tu_ngay:%d/%m/%Y}–{p.den_ngay:%d/%m/%Y}"
    return f"Hoàn vật tư ứng cho thợ {p.tho.ho_ten}, đối soát {p.tu_ngay:%d/%m/%Y}–{p.den_ngay:%d/%m/%Y}"


# (tiêu đề cột, hàm lấy giá trị từ phiếu). Tài khoản nợ/có để trống (OI-07).
COT_MISA: list[tuple[str, callable]] = [
    ("Ngày chứng từ", lambda p: timezone.localdate(p.xac_nhan_luc).strftime("%d/%m/%Y")),
    ("Số chứng từ", lambda p: p.ma),
    ("Diễn giải", _dien_giai),
    ("Đối tượng", lambda p: p.tho.ma_tho),
    ("Tên đối tượng", lambda p: p.tho.ho_ten),
    ("Số tiền", lambda p: p.so_phai_nop),
    ("TK Nợ", lambda p: ""),
    ("TK Có", lambda p: ""),
]


def phieu_trong_ky(nguoi_dung, tu_ngay: date, den_ngay: date) -> list[PhieuDoiSoat]:
    kiem_tra_vai_tro(nguoi_dung, *XEM_DOI_SOAT)
    return list(
        PhieuDoiSoat.objects.filter(
            trang_thai=PhieuDoiSoat.TrangThai.DA_NOP,
            xac_nhan_luc__date__gte=tu_ngay,
            xac_nhan_luc__date__lte=den_ngay,
        )
        .select_related("tho")
        .order_by("xac_nhan_luc", "id")
    )


def bang_xuat(ds_phieu) -> list[list]:
    return [[ten for ten, _ in COT_MISA]] + [[lay(p) for _, lay in COT_MISA] for p in ds_phieu]


def xuat_csv(ds_phieu) -> bytes:
    """UTF-8 có BOM để Excel/MISA mở đúng tiếng Việt."""
    dem = io.StringIO()
    csv.writer(dem).writerows(bang_xuat(ds_phieu))
    return ("﻿" + dem.getvalue()).encode("utf-8")


def xuat_xlsx(ds_phieu) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "DoiSoat"
    for dong in bang_xuat(ds_phieu):
        ws.append(dong)
    for o in ws[1]:
        o.font = Font(bold=True)
    for o in ws["F"][1:]:
        o.number_format = "#,##0"
    for cot, rong in zip("ABCDEFGH", (14, 20, 60, 10, 26, 14, 8, 8), strict=True):
        ws.column_dimensions[cot].width = rong
    dem = io.BytesIO()
    wb.save(dem)
    return dem.getvalue()
