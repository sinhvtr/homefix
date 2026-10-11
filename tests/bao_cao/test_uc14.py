"""UC14 — báo cáo vận hành và doanh thu: số liệu khớp bộ dữ liệu kiểm thử cố định (tuần 01–07/09/2026)."""

import io
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from django.core.exceptions import PermissionDenied
from openpyxl import load_workbook

from apps.bao_cao.dao import BaoCaoDAO, chon_ky
from apps.bao_cao.views import phan_tram, phut
from apps.don_sua_chua.models import DonSuaChua
from apps.he_thong.nhat_ky import ghi_nhat_ky
from apps.tai_chinh.models import ThanhToan
from apps.tai_khoan.models import VaiTro
from apps.tiep_nhan_dieu_phoi.models import PhanCong
from tests.conftest import DiaChiFactory, DichVuFactory, KhuVucFactory, NguoiDungFactory, ThoFactory, tao_don

pytestmark = pytest.mark.django_db
VN = ZoneInfo("Asia/Ho_Chi_Minh")
KY = chon_ky("tuy_chon", date(2026, 9, 1), date(2026, 9, 7))


def _luc(ngay, gio, phut_=0):
    """Thời điểm tháng 9/2026, giờ Việt Nam."""
    return datetime(2026, 9, ngay, gio, phut_, tzinfo=VN)


@pytest.fixture
def du_lieu():
    """Bộ dữ liệu cố định — xem docstring từng đơn để đối chiếu kết quả mong đợi."""
    td = NguoiDungFactory(vai_tro=VaiTro.TONG_DAI)
    tho_a, tho_b = ThoFactory(ho_ten="An"), ThoFactory(ho_ten="Bình")
    kv1, kv2 = KhuVucFactory(ten="Ba Đình"), KhuVucFactory(ten="Cầu Giấy")
    vs, sc = (
        DichVuFactory(ten_dich_vu="Vệ sinh"),
        DichVuFactory(ten_dich_vu="Sửa chữa", nhom_dich_vu="SUA_CHUA_THAY_VAT_TU"),
    )

    def don(trang_thai, nhan_luc, dv=vs, kv=kv1, **kw):
        d = tao_don(trang_thai, dia_chi=DiaChiFactory(khu_vuc=kv), dich_vu=dv, nguoi_tao=td, **kw)
        DonSuaChua.objects.filter(pk=d.pk).update(ngay_tiep_nhan=nhan_luc)
        return DonSuaChua.objects.get(pk=d.pk)

    def xong(d, tho, luc, tien, hinh_thuc="TIEN_MAT", trang_thai="CHUA_NOP"):
        DonSuaChua.objects.filter(pk=d.pk).update(tho=tho, tong_tien_phai_thu=tien)
        ThanhToan.objects.create(
            don=d, tho=tho, so_tien=tien, hinh_thuc=hinh_thuc, trang_thai=trang_thai, thoi_diem_thu=luc
        )

    def pc(d, tho, giao, trang_thai, phan_hoi=None):
        PhanCong.objects.create(
            don=d,
            tho=tho,
            nguoi_giao=td,
            thoi_diem_giao=giao,
            han_phan_hoi=giao + timedelta(minutes=15),
            trang_thai=trang_thai,
            thoi_diem_phan_hoi=phan_hoi,
            tien_cong_du_kien=150_000,
        )

    # d0: tiếp nhận NGOÀI kỳ, hoàn thành TRONG kỳ (02/09) — 200k, thợ An, Ba Đình, CK thợ đã nộp
    d0 = don("HOAN_THANH", datetime(2026, 8, 25, 9, 0, tzinfo=VN))
    xong(d0, tho_a, _luc(2, 16), 200_000, "CK_TAI_KHOAN_THO", "DA_NOP")
    # d1: nhận 01/09 08:00, An nhận việc 08:20 (20 phút), xong 01/09 — 300k tiền mặt chưa nộp
    d1 = don("HOAN_THANH", _luc(1, 8))
    pc(d1, tho_a, _luc(1, 8, 5), "DA_NHAN", _luc(1, 8, 20))
    xong(d1, tho_a, _luc(1, 15), 300_000)
    # d2: nhận 02/09 09:00, Bình nhận 10:00 (60 phút), xong 03/09 — 500k sửa chữa, Cầu Giấy, CK công ty
    d2 = don("HOAN_THANH", _luc(2, 9), dv=sc, kv=kv2)
    pc(d2, tho_b, _luc(2, 9, 30), "DA_NHAN", _luc(2, 10))
    xong(d2, tho_b, _luc(3, 11), 500_000, "CK_TAI_KHOAN_CONG_TY", "KHONG_CAN_NOP")
    # d3: nhận 03/09, An từ chối, Bình quá hạn → gắn cờ chưa cử được, khách hủy khi đang chờ cử thợ
    d3 = don("DA_HUY", _luc(3, 10), ly_do_huy="Không có thợ")
    pc(d3, tho_a, _luc(3, 10, 5), "TU_CHOI", _luc(3, 10, 6))
    pc(d3, tho_b, _luc(3, 10, 5), "QUA_HAN")
    ghi_nhat_ky(td, "CHUA_CU_DUOC_THO", d3, du_lieu_sau={"ly_do": "Không có thợ"})
    ghi_nhat_ky(
        td, "DOI_TRANG_THAI", d3, {"trang_thai": "CHO_CU_THO"}, {"trang_thai": "DA_HUY", "chuyen": "T14"}
    )
    # d4: nhận 04/09, khách hủy khi còn chờ xác nhận lịch (không tính "không cử được")
    d4 = don("DA_HUY", _luc(4, 9), ly_do_huy="Khách tự sửa được")
    ghi_nhat_ky(
        td,
        "DOI_TRANG_THAI",
        d4,
        {"trang_thai": "CHO_XAC_NHAN_LICH"},
        {"trang_thai": "DA_HUY", "chuyen": "T14"},
    )
    # d5: nhận 05/09 08:00, An nhận 08:10 (10 phút), đang sửa
    d5 = don("DANG_SUA", _luc(5, 8))
    pc(d5, tho_a, _luc(5, 8, 2), "DA_NHAN", _luc(5, 8, 10))
    # bh: đơn bảo hành tạo 06/09 cho d0 (thợ An làm đơn gốc)
    don("CHO_CU_THO", _luc(6, 9), loai_don="BAO_HANH", don_goc=d0)
    # d9: ngoài kỳ hoàn toàn — chỉ góp vào công nợ hiện tại (100k chưa nộp của Bình)
    d9 = don("HOAN_THANH", _luc(10, 9))
    xong(d9, tho_b, _luc(10, 15), 100_000)
    return {"tho_a": tho_a, "tho_b": tho_b}


def test_uc14_so_lieu_khop_bo_du_lieu_co_dinh(du_lieu):
    bc = BaoCaoDAO.lap(NguoiDungFactory(vai_tro=VaiTro.GIAM_DOC), KY)
    assert (bc.so_tiep_nhan, bc.so_hoan_thanh, bc.so_huy) == (6, 3, 2)
    assert dict(bc.huy_theo_ly_do) == {"Không có thợ": 1, "Khách tự sửa được": 1}
    # Không cử được: chỉ d3 (cờ + hủy khi đang chờ cử thợ) / 6 đơn tiếp nhận
    assert (bc.khong_cu_duoc.tu_so, bc.khong_cu_duoc.mau_so) == (1, 6)
    # Tiếp nhận → có thợ: 20, 60, 10 phút → trung vị 20; 2/3 có thợ trong 30 phút
    assert bc.trung_vi_phut_co_tho == 20
    assert (bc.co_tho_trong_30_phut.tu_so, bc.co_tho_trong_30_phut.mau_so) == (2, 3)
    # Doanh thu = Σ tổng phải thu của đơn hoàn thành trong kỳ (theo lúc thu tiền — A-85)
    assert bc.doanh_thu == 1_000_000
    assert dict((n, t) for n, t in bc.doanh_thu_theo_ngay if t) == {
        date(2026, 9, 1): 300_000,
        date(2026, 9, 2): 200_000,
        date(2026, 9, 3): 500_000,
    }
    assert len(bc.doanh_thu_theo_ngay) == 7
    assert set(bc.doanh_thu_theo_dich_vu) == {("Vệ sinh", 2, 500_000), ("Sửa chữa", 1, 500_000)}
    assert set(bc.doanh_thu_theo_tho) == {("An", 2, 500_000), ("Bình", 1, 500_000)}
    assert set(bc.doanh_thu_theo_khu_vuc) == {("Ba Đình", 2, 500_000), ("Cầu Giấy", 1, 500_000)}
    # Tỷ lệ làm lại (A-84) = 1 đơn bảo hành / 3 đơn thường hoàn thành; theo thợ làm đơn gốc
    assert (bc.lam_lai.tu_so, bc.lam_lai.mau_so) == (1, 3)
    assert {(d.ten, d.tu_so, d.mau_so) for d in bc.lam_lai_theo_tho} == {("An", 1, 2), ("Bình", 0, 1)}
    # Công nợ thợ hiện tại (không phụ thuộc kỳ): 300k của An + 100k của Bình
    assert bc.cong_no_hien_tai == 400_000 and bc.cong_no_theo_tho == [("An", 300_000), ("Bình", 100_000)]
    # Đề nghị gửi trong kỳ theo thợ
    de_nghi = {d["tho"]: (d["tong"], d["nhan"], d["tu_choi"], d["qua_han"]) for d in bc.de_nghi_theo_tho}
    assert de_nghi == {"An": (3, 2, 1, 0), "Bình": (2, 1, 0, 1)}
    don_ngay = {n: (a, b) for n, a, b in bc.don_theo_ngay}
    assert don_ngay[date(2026, 9, 1)] == (1, 1) and don_ngay[date(2026, 9, 3)] == (1, 1)
    assert don_ngay[date(2026, 9, 6)] == (1, 0) and sum(a for a, _ in don_ngay.values()) == 6


def test_uc14_ke_toan_chi_thay_tai_chinh(du_lieu, client):
    kt = NguoiDungFactory(vai_tro=VaiTro.KE_TOAN)
    bc = BaoCaoDAO.lap(kt, KY)
    assert not bc.xem_van_hanh and bc.khong_cu_duoc is None and bc.doanh_thu == 1_000_000
    client.force_login(kt)
    trang = client.get(
        "/bao-cao/", {"ky": "tuy_chon", "tu": "2026-09-01", "den": "2026-09-07"}
    ).content.decode()
    assert "1.000.000 đ" in trang and "Tỷ lệ làm lại" not in trang and "Đề nghị nhận việc" not in trang


@pytest.mark.parametrize("vai_tro", [VaiTro.TONG_DAI, VaiTro.THO])
def test_uc14_khong_co_quyen(vai_tro):
    nguoi = ThoFactory().nguoi_dung if vai_tro == VaiTro.THO else NguoiDungFactory(vai_tro=vai_tro)
    with pytest.raises(PermissionDenied):
        BaoCaoDAO.lap(nguoi, KY)


def test_uc14_man_hinh_giam_doc(du_lieu, client):
    client.force_login(NguoiDungFactory(vai_tro=VaiTro.GIAM_DOC))
    trang = client.get(
        "/bao-cao/", {"ky": "tuy_chon", "tu": "2026-09-01", "den": "2026-09-07"}
    ).content.decode()
    for chuoi in (
        "01/09/2026 – 07/09/2026",
        "1.000.000 đ",
        "16,7%",
        "20 phút",
        "33,3%",
        "Không có thợ",
        "chart.umd.min.js",
    ):
        assert chuoi in trang, chuoi
    assert 'id="du-lieu-bieu-do"' in trang


def test_uc14_xuat_xlsx_moi_bang(du_lieu, client):
    client.force_login(NguoiDungFactory(vai_tro=VaiTro.QUAN_TRI))
    phan_hoi = client.get("/bao-cao/xuat/", {"ky": "tuy_chon", "tu": "2026-09-01", "den": "2026-09-07"})
    assert phan_hoi["Content-Disposition"] == 'attachment; filename="bao_cao_20260901_20260907.xlsx"'
    wb = load_workbook(io.BytesIO(phan_hoi.content))
    assert wb.sheetnames == [
        "TongQuan",
        "DonTheoNgay",
        "DoanhThuTheoNgay",
        "DoanhThuTheoDichVu",
        "DoanhThuTheoTho",
        "DoanhThuTheoKhuVuc",
        "CongNoTho",
        "HuyTheoLyDo",
        "LamLaiTheoTho",
        "DeNghiTheoTho",
    ]
    tong = {r[0].value: r[1].value for r in wb["TongQuan"].iter_rows()}
    assert tong["Doanh thu (đ)"] == 1_000_000 and tong["Số đơn tiếp nhận"] == 6
    assert abs(tong["Tỷ lệ làm lại"] - 1 / 3) < 1e-9
    assert wb["DoanhThuTheoNgay"].max_row == 8  # tiêu đề + 7 ngày
    client.force_login(NguoiDungFactory(vai_tro=VaiTro.KE_TOAN))
    wb_kt = load_workbook(io.BytesIO(client.get("/bao-cao/xuat/", {"ky": "7_ngay"}).content))
    assert "LamLaiTheoTho" not in wb_kt.sheetnames and "CongNoTho" in wb_kt.sheetnames


@pytest.mark.parametrize(
    ("ma", "tu", "den", "mong_doi"),
    [
        ("hom_nay", None, None, (date(2026, 10, 11), date(2026, 10, 11))),
        ("7_ngay", None, None, (date(2026, 10, 5), date(2026, 10, 11))),
        ("thang_nay", None, None, (date(2026, 10, 1), date(2026, 10, 11))),
        ("tuy_chon", date(2026, 9, 9), date(2026, 9, 1), (date(2026, 9, 1), date(2026, 9, 9))),
        ("tuy_chon", None, None, (date(2026, 10, 5), date(2026, 10, 11))),  # thiếu ngày → mặc định 7 ngày
        ("la", None, None, (date(2026, 10, 5), date(2026, 10, 11))),
    ],
)
def test_uc14_bo_chon_ky(ma, tu, den, mong_doi):
    ky = chon_ky(ma, tu, den, hom_nay=date(2026, 10, 11))
    assert (ky.tu, ky.den) == mong_doi
    assert ky.ket_thuc - ky.bat_dau == timedelta(days=(ky.den - ky.tu).days + 1)
    assert ky.bat_dau.utcoffset() == timedelta(hours=7)  # mốc ngày theo giờ Việt Nam


def test_uc14_ky_trong_khong_loi(client):
    client.force_login(NguoiDungFactory(vai_tro=VaiTro.GIAM_DOC))
    trang = client.get("/bao-cao/", {"ky": "hom_nay"})
    assert trang.status_code == 200 and "Không có đơn hoàn thành trong kỳ." in trang.content.decode()


@pytest.mark.parametrize(("vao", "ra"), [(None, "—"), (0.1666, "16,7%"), (1, "100,0%")])
def test_dinh_dang_phan_tram(vao, ra):
    assert phan_tram(vao) == ra


@pytest.mark.parametrize(("vao", "ra"), [(None, "—"), (20, "20 phút"), (150, "2,5 giờ")])
def test_dinh_dang_phut(vao, ra):
    assert phut(vao) == ra
