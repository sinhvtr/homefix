"""Xuất XLSX mọi bảng của báo cáo (UC14) — mỗi bảng một trang tính. Kế toán chỉ nhận các bảng tài chính."""

import io

from openpyxl import Workbook
from openpyxl.styles import Font

from .dao import BaoCao

DINH_DANG_TIEN = "#,##0"
DINH_DANG_TY_LE = "0.0%"


def _trang(wb, ten, tieu_de, dong, cot_tien=(), cot_ty_le=()):
    ws = wb.create_sheet(ten[:31])
    ws.append(tieu_de)
    for o in ws[1]:
        o.font = Font(bold=True)
    for d in dong:
        ws.append(list(d))
    for cot, dinh_dang in [(c, DINH_DANG_TIEN) for c in cot_tien] + [(c, DINH_DANG_TY_LE) for c in cot_ty_le]:
        for (o,) in ws.iter_rows(min_row=2, min_col=cot, max_col=cot):
            o.number_format = dinh_dang
    for i, _ in enumerate(tieu_de, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = 18 if i > 1 else 32
    return ws


def xuat_xlsx(bc: BaoCao) -> bytes:
    wb = Workbook()
    tong = wb.active
    tong.title = "TongQuan"
    tong.append(["Kỳ báo cáo", bc.ky.nhan])
    tong.append(["Số đơn tiếp nhận", bc.so_tiep_nhan])
    tong.append(["Số đơn hoàn thành", bc.so_hoan_thanh])
    tong.append(["Doanh thu (đ)", bc.doanh_thu])
    tong.append(["Công nợ thợ hiện tại (đ)", bc.cong_no_hien_tai])
    if bc.xem_van_hanh:
        tong.append(["Số đơn hủy", bc.so_huy])
        tong.append(["Tỷ lệ không cử được thợ", bc.khong_cu_duoc.ty_le])
        tong.append(["Trung vị tiếp nhận → có thợ (phút)", bc.trung_vi_phut_co_tho])
        tong.append([bc.co_tho_trong_30_phut.ten, bc.co_tho_trong_30_phut.ty_le])
        tong.append(["Tỷ lệ làm lại", bc.lam_lai.ty_le])
    tong.column_dimensions["A"].width = 36
    tong.column_dimensions["B"].width = 26
    for o in tong["B"][3:5]:
        o.number_format = DINH_DANG_TIEN

    _trang(
        wb,
        "DonTheoNgay",
        ["Ngày", "Tiếp nhận", "Hoàn thành"],
        ((f"{n:%d/%m/%Y}", a, b) for n, a, b in bc.don_theo_ngay),
    )
    _trang(
        wb,
        "DoanhThuTheoNgay",
        ["Ngày", "Doanh thu (đ)"],
        ((f"{n:%d/%m/%Y}", t) for n, t in bc.doanh_thu_theo_ngay),
        [2],
    )
    for ten, bang in (
        ("DoanhThuTheoDichVu", bc.doanh_thu_theo_dich_vu),
        ("DoanhThuTheoTho", bc.doanh_thu_theo_tho),
        ("DoanhThuTheoKhuVuc", bc.doanh_thu_theo_khu_vuc),
    ):
        _trang(wb, ten, [ten.replace("DoanhThuTheo", ""), "Số đơn", "Doanh thu (đ)"], bang, [3])
    _trang(wb, "CongNoTho", ["Thợ", "Chưa nộp (đ)"], bc.cong_no_theo_tho, [2])
    if bc.xem_van_hanh:
        _trang(wb, "HuyTheoLyDo", ["Lý do hủy", "Số đơn"], bc.huy_theo_ly_do)
        _trang(
            wb,
            "LamLaiTheoTho",
            ["Thợ", "Đơn bảo hành", "Đơn hoàn thành", "Tỷ lệ"],
            ((d.ten, d.tu_so, d.mau_so, d.ty_le) for d in bc.lam_lai_theo_tho),
            cot_ty_le=[4],
        )
        _trang(
            wb,
            "DeNghiTheoTho",
            ["Thợ", "Đề nghị", "Nhận", "Từ chối", "Quá hạn", "Đã hủy", "Tỷ lệ nhận"],
            (
                (d["tho"], d["tong"], d["nhan"], d["tu_choi"], d["qua_han"], d["huy"], d["ty_le_nhan"])
                for d in bc.de_nghi_theo_tho
            ),
            cot_ty_le=[7],
        )
    dem = io.BytesIO()
    wb.save(dem)
    return dem.getvalue()
